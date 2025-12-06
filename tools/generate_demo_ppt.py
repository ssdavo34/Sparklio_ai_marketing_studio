import os
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    print("python-pptx not found. Installing...")
    os.system("pip install python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Helper to set title and content
    def add_slide(title_text, content_text_list):
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        title.text = title_text
        
        # Customize Title
        for paragraph in title.text_frame.paragraphs:
            paragraph.font.size = Pt(40)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(59, 130, 246) # Blue accent

        content = slide.placeholders[1]
        tf = content.text_frame
        tf.text = content_text_list[0]
        
        for i in range(1, len(content_text_list)):
            p = tf.add_paragraph()
            p.text = content_text_list[i]
            p.font.size = Pt(20)

    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Sparklio AI Marketing Studio"
    subtitle.text = "Phase 10 완료 보고 및 데모\n\n[팀 구성]\nA팀: QA (품질 보증)\nB팀: Backend (서버 개발)\nC팀: 박성언 (PM 및 총괄)"

    # Customize Title Slide
    title.text_frame.paragraphs[0].font.size = Pt(54)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(15, 23, 42) # Dark Blue

    # 2. Project Overview
    add_slide("프로젝트 개요", [
        "목표: AI 기반 마케팅 콘텐츠 생성 및 편집 플랫폼 구축",
        "현재 상태: Phase 10 완료 (전체 진행률 67%)",
        "주요 성과:",
        "  - 하이브리드 인프라 (Mac Mini + Desktop GPU) 구축 완료",
        "  - Multi-LLM Gateway (OpenAI, Gemini, Ollama) 통합",
        "  - Spark Chat 및 Canvas Studio UI/UX 고도화"
    ])

    # 3. System Architecture
    add_slide("시스템 아키텍처", [
        "하이브리드 클라우드/온프레미스 구조",
        "",
        "1. Control Tower (Mac Mini)",
        "  - FastAPI Backend Server",
        "  - Next.js Frontend Server",
        "  - Redis (Caching) & PostgreSQL (DB)",
        "",
        "2. GPU Worker (Desktop PC)",
        "  - Ollama (Llama 3, Mistral 등 로컬 LLM)",
        "  - ComfyUI (이미지 생성 엔진)",
        "  - NVIDIA RTX 4090 활용"
    ])

    # 4. Spark Chat (Key Feature 1)
    add_slide("핵심 기능 1: Spark Chat", [
        "지능형 멀티 모델 채팅 인터페이스",
        "",
        "주요 특징:",
        "  - 다양한 LLM 선택 가능 (GPT-4o, Gemini 2.0, Claude 3.5)",
        "  - Live Mode: 실제 API 연동 완료",
        "  - UI 개선: 하단 고정 입력창 (VS Code 스타일)",
        "  - 컨텍스트 인식: 이전 대화 내용을 바탕으로 답변"
    ])

    # 5. Canvas Studio (Key Feature 2)
    add_slide("핵심 기능 2: Canvas Studio", [
        "Konva.js 기반 이미지 편집 스튜디오",
        "",
        "주요 특징:",
        "  - 레이어 기반 편집 (Photoshop 유사 경험)",
        "  - AI 생성 이미지의 실시간 수정 및 배치",
        "  - 브랜드 키트 연동 (로고, 색상 팔레트)",
        "  - 직관적인 속성 패널 (크기, 색상, 폰트 조절)"
    ])

    # 6. Future Roadmap
    add_slide("향후 계획 (Phase 11-15)", [
        "남은 과제 및 로드맵",
        "",
        "Phase 11: LLM Router 고도화 (자동 모델 추천)",
        "Phase 12: Canvas Drag & Drop 및 Undo/Redo 구현",
        "Phase 13: ComfyUI 이미지 생성 엔진 완전 통합",
        "Phase 14: 사용자 인증 및 권한 관리 시스템",
        "Phase 15: 최종 통합 테스트 및 배포"
    ])

    output_path = "sparklio_demo.pptx"
    prs.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == "__main__":
    create_presentation()
